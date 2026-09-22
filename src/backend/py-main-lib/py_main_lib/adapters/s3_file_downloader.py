"""Загрузка и выгрузка файлов из/в S3."""

import asyncio
import base64
import io
import mimetypes
import time
import uuid
from typing import Any, cast

import aioboto3
import httpx
from botocore.exceptions import BotoCoreError, ClientError, EndpointConnectionError

from py_main_lib.colored_logger import get_colored_logger


try:
    from docx import Document

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import PyPDF2

    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


logger = get_colored_logger(__name__)


class S3FileDownloader:
    """Загрузка файлов из S3."""

    def __init__(
        self,
        base_s3_url: str | None = None,
        system_token: str | None = None,
        *,
        bucket_name: str | None = None,
        temp_prefix: str = "temp-files/rag-indexer/",
        retry_attempts: int = 3,
        client_params: dict[str, Any] | None = None,
        session: aioboto3.Session | None = None,
    ):
        """Инициализация с базовым URL S3 API или прямым доступом к бакету."""
        self.base_s3_url = base_s3_url.rstrip("/") if base_s3_url else None
        self.system_token = system_token
        self.bucket_name = bucket_name
        self.temp_prefix = self._normalize_prefix(temp_prefix)
        self.retry_attempts = max(1, retry_attempts)
        self.client_params = client_params or {}
        self._session = session or aioboto3.Session()

    def _ensure_bucket_configured(self) -> None:
        if not self.bucket_name:
            raise ValueError("bucket_name is required for direct S3 operations")

    async def _get_object_bytes(self, s3_key: str) -> bytes:
        """Скачивает объект из S3 с retry и возвратом байтов."""
        self._ensure_bucket_configured()
        normalized_key = self._normalize_key(s3_key)

        attempt = 0
        delay = 1.0
        last_error: Exception | None = None

        while attempt < self.retry_attempts:
            attempt += 1
            start_time = time.perf_counter()
            try:
                async with self._session.client("s3", **self.client_params) as client:
                    response = await client.get_object(Bucket=self.bucket_name, Key=normalized_key)
                    body: bytes = await response["Body"].read()

                elapsed = time.perf_counter() - start_time
                logger.info(
                    f"Downloaded s3://{self.bucket_name}/{normalized_key} "
                    f"({len(body)} bytes) in {elapsed:.3f}s (attempt {attempt})"
                )
                return body

            except (ClientError, EndpointConnectionError, BotoCoreError) as exc:
                last_error = exc
                logger.warning(
                    f"S3 download attempt {attempt}/{self.retry_attempts} failed for "
                    f"{normalized_key}: {type(exc).__name__}: {exc}"
                )
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                logger.error(
                    f"Unexpected error downloading {normalized_key} on attempt "
                    f"{attempt}/{self.retry_attempts}: {type(exc).__name__}: {exc}"
                )

            if attempt < self.retry_attempts:
                await asyncio.sleep(delay)
                delay *= 2

        raise RuntimeError(
            f"Failed to download s3://{self.bucket_name}/{normalized_key} " f"after {self.retry_attempts} attempts"
        ) from last_error

    async def download_text_file(self, s3_key: str) -> str:
        """Скачивание текстового файла из S3 temp-bucket с fallback по кодировкам."""
        content_bytes = await self._get_object_bytes(s3_key)
        for encoding in ("utf-8", "cp1251", "latin-1"):
            try:
                decoded = content_bytes.decode(encoding)
                used_encoding = encoding
                break
            except UnicodeDecodeError:
                continue
        else:
            decoded = content_bytes.decode("utf-8", errors="replace")
            used_encoding = "utf-8 (replace)"

        logger.info(
            f"Decoded text file {self._normalize_key(s3_key)} with encoding {used_encoding} "
            f"({len(content_bytes)} bytes)"
        )
        return decoded

    async def delete_temp_file(self, s3_key: str) -> bool:
        """Удаляет временный файл из S3 с retry логикой."""
        self._ensure_bucket_configured()
        normalized_key = self._normalize_key(s3_key)

        attempt = 0
        delay = 1.0
        last_error: Exception | None = None

        while attempt < self.retry_attempts:
            attempt += 1
            start_time = time.perf_counter()
            try:
                async with self._session.client("s3", **self.client_params) as client:
                    await client.delete_object(Bucket=self.bucket_name, Key=normalized_key)

                elapsed = time.perf_counter() - start_time
                logger.info(
                    f"Deleted temp file s3://{self.bucket_name}/{normalized_key} "
                    f"in {elapsed:.3f}s (attempt {attempt})"
                )
                return True

            except (ClientError, EndpointConnectionError, BotoCoreError) as exc:
                last_error = exc
                logger.warning(
                    f"S3 delete attempt {attempt}/{self.retry_attempts} failed for "
                    f"{normalized_key}: {type(exc).__name__}: {exc}"
                )
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                logger.error(
                    f"Unexpected error deleting {normalized_key} on attempt "
                    f"{attempt}/{self.retry_attempts}: {type(exc).__name__}: {exc}"
                )

            if attempt < self.retry_attempts:
                await asyncio.sleep(delay)
                delay *= 2

        logger.error(
            f"Failed to delete temp file s3://{self.bucket_name}/{normalized_key} "
            f"after {self.retry_attempts} attempts: {last_error}"
        )
        return False

    async def get_file_download_link(self, file_key: str) -> str:
        """Получение ссылки для скачивания файла из S3."""
        if not self.base_s3_url:
            raise ValueError("base_s3_url is required for HTTP download link requests")

        url = f"{self.base_s3_url}/v1/files/download_link"

        # Подготовка заголовков
        headers: dict[str, str] = {"Content-Type": "application/json"}
        if self.system_token:
            headers["Authorization"] = f"Bearer {self.system_token}"
            headers["X-Service-Name"] = "aitunnel-adapter"
            logger.info(f"Requesting download link for {file_key} from {url} with system token")
        else:
            logger.info(f"Requesting download link for {file_key} from {url} without system token")

        async with httpx.AsyncClient(timeout=30.0) as client:
            try:
                response = await client.post(url, json={"file_key": file_key}, headers=headers)
                response.raise_for_status()

                data = cast(dict[str, Any], response.json())
                payload = data.get("payload")
                if isinstance(payload, dict):
                    download_url = payload.get("file_download_link")
                else:
                    download_url = None

                if not isinstance(download_url, str) or not download_url:
                    raise ValueError(f"No download link received for file: {file_key}")

                logger.info(f"Got download link for {file_key}: {download_url}")
                return download_url

            except httpx.HTTPStatusError as e:
                status_code = e.response.status_code
                response_text = e.response.text if hasattr(e.response, "text") else str(e)
                logger.error(f"HTTP {status_code} error getting download link for {file_key}: {response_text}")

                if status_code == 404:
                    raise Exception(f"Файл '{file_key}' не найден на S3 сервере")
                elif status_code == 403:
                    raise Exception(f"Нет прав доступа к файлу '{file_key}'")
                elif status_code == 500:
                    raise Exception(f"S3 сервер недоступен для файла '{file_key}'")
                else:
                    raise Exception(f"HTTP {status_code} ошибка при получении ссылки на '{file_key}': {e}")
            except Exception as e:
                logger.error(f"Unexpected error getting download link for {file_key}: {type(e).__name__}: {e}")
                raise Exception(f"Не удалось получить ссылку на '{file_key}': {str(e)}")

    async def download_file_as_base64(self, file_key: str) -> tuple[str, str]:
        """Загрузка файла и конвертация в base64."""
        try:
            download_url = await self.get_file_download_link(file_key)
            logger.info(f"Downloading file from URL: {download_url}")

            async with httpx.AsyncClient(timeout=60.0) as client:
                response = await client.get(download_url)
                response.raise_for_status()

                file_content = response.content
                mime_type = self._get_mime_type(file_key)
                base64_content = base64.b64encode(file_content).decode("utf-8")

                logger.info(f"Downloaded {file_key} as base64 ({mime_type}, {len(file_content)} bytes)")
                return base64_content, mime_type

        except Exception as e:
            logger.error(f"Error downloading file {file_key}: {e}")
            raise

    def _get_mime_type(self, file_key: str) -> str:
        """Определение MIME типа по расширению файла."""
        mime_type, _ = mimetypes.guess_type(file_key)

        if mime_type is None:
            extension = file_key.lower().split(".")[-1]
            if extension in ["jpg", "jpeg"]:
                return "image/jpeg"
            elif extension == "png":
                return "image/png"
            elif extension == "gif":
                return "image/gif"
            elif extension == "webp":
                return "image/webp"
            elif extension == "pdf":
                return "application/pdf"
            else:
                return "application/octet-stream"

        return mime_type

    def is_image_file(self, file_key: str) -> bool:
        """Проверка, является ли файл изображением."""
        mime_type = self._get_mime_type(file_key)
        return mime_type.startswith("image/")

    async def upload_base64_image(self, base64_data: str, user_id: str, extension: str = "png") -> str:
        """Загрузка base64 изображения на S3 и получение file_key."""
        url = f"{self.base_s3_url}/v1/files/upload_by_system_user"

        file_name = f"generated_image_{uuid.uuid4().hex}.{extension}"

        try:
            image_bytes = base64.b64decode(base64_data)
        except Exception as e:
            raise ValueError(f"Invalid base64 data: {e}")

        headers: dict[str, str] = {}
        if self.system_token:
            headers["Authorization"] = f"Bearer {self.system_token}"
            headers["X-Service-Name"] = "aitunnel-adapter"
            logger.info(f"Uploading image {file_name} to {url} with system token")
        else:
            logger.info(f"Uploading image {file_name} to {url} without system token")

        files = {"file": (file_name, io.BytesIO(image_bytes), f"image/{extension}")}

        params = {"is_library": False, "user_id": user_id}

        async with httpx.AsyncClient(timeout=60.0) as client:
            try:
                response = await client.post(url, files=files, params=params, headers=headers)
                response.raise_for_status()

                result = cast(dict[str, Any], response.json())
                payload = result.get("payload")
                file_key = payload.get("file_key") if isinstance(payload, dict) else None

                if not isinstance(file_key, str) or not file_key:
                    raise ValueError("No file_key received for uploaded image")

                logger.info(f"Successfully uploaded image as {file_key}")
                return file_key

            except httpx.HTTPStatusError as e:
                status_code = e.response.status_code
                response_text = e.response.text if hasattr(e.response, "text") else str(e)
                logger.error(f"HTTP {status_code} error uploading image: {response_text}")

                if status_code == 403:
                    raise Exception("Нет прав для загрузки файлов на S3")
                elif status_code == 413:
                    raise Exception("Изображение слишком большое для загрузки")
                elif status_code == 500:
                    raise Exception("S3 сервер недоступен для загрузки")
                else:
                    raise Exception(f"HTTP {status_code} ошибка при загрузке изображения: {e}")
            except Exception as e:
                logger.error(f"Unexpected error uploading image: {type(e).__name__}: {e}")
                raise Exception(f"Не удалось загрузить изображение на S3: {str(e)}")

    def is_docx_file(self, file_key: str) -> bool:
        """Проверка, является ли файл DOCX документом."""
        return file_key.lower().endswith(".docx")

    async def download_and_extract_docx_text(self, file_key: str) -> str:
        """Скачивает DOCX файл и извлекает из него текст."""
        if not DOCX_AVAILABLE:
            raise ImportError(
                "python-docx library is required for DOCX processing. Install with: pip install python-docx"
            )

        if not self.is_docx_file(file_key):
            raise ValueError(f"File {file_key} is not a DOCX document")

        try:
            download_url = await self.get_file_download_link(file_key)
            logger.info(f"Downloading DOCX file from URL: {download_url}")

            async with httpx.AsyncClient(timeout=60.0) as client:
                response = await client.get(download_url)
                response.raise_for_status()

                # Создаем Document из байтов
                doc = Document(io.BytesIO(response.content))

                # Извлекаем весь текст из документа
                full_text = []
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        full_text.append(paragraph.text.strip())

                # Извлекаем текст из таблиц
                for table in doc.tables:
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            full_text.append(" | ".join(row_text))

                extracted_text = "\n".join(full_text)
                logger.info(f"Extracted {len(extracted_text)} characters from DOCX: {file_key}")

                return extracted_text

        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_key}: {e}")
            raise

    def is_txt_file(self, file_key: str) -> bool:
        """Проверка, является ли файл текстовым документом."""
        txt_extensions = [".txt", ".text", ".md", ".csv", ".json", ".xml", ".log"]
        return any(file_key.lower().endswith(ext) for ext in txt_extensions)

    async def download_and_extract_txt_text(self, file_key: str) -> str:
        """Скачивает текстовый файл и извлекает из него содержимое."""
        if not self.is_txt_file(file_key):
            raise ValueError(f"File {file_key} is not a text document")

        try:
            download_url = await self.get_file_download_link(file_key)
            logger.info(f"Downloading text file from URL: {download_url}")

            async with httpx.AsyncClient(timeout=60.0) as client:
                response = await client.get(download_url)
                response.raise_for_status()

                # Пытаемся декодировать в UTF-8, если не получается - пробуем другие кодировки
                decoded_text: str
                try:
                    decoded_text = response.content.decode("utf-8")
                except UnicodeDecodeError:
                    try:
                        decoded_text = response.content.decode("cp1251")
                    except UnicodeDecodeError:
                        decoded_text = response.content.decode("latin-1")

                # Очищаем и проверяем содержимое
                text_content = decoded_text.strip()

                if not text_content:
                    logger.warning(f"Text file {file_key} is empty")
                    return ""

                logger.info(f"Extracted {len(text_content)} characters from text file: {file_key}")
                return text_content

        except Exception as e:
            logger.error(f"Error extracting text from file {file_key}: {e}")
            raise Exception(f"Не удалось извлечь текст из файла: {str(e)}")

    def is_pdf_file(self, file_key: str) -> bool:
        """Проверка, является ли файл PDF документом."""
        return file_key.lower().endswith(".pdf")

    async def download_and_extract_pdf_text(self, file_key: str) -> str:
        """Скачивает PDF файл и извлекает из него текст."""
        if not PDF_AVAILABLE:
            raise ImportError("PyPDF2 library is required for PDF processing. Install with: pip install PyPDF2")

        if not self.is_pdf_file(file_key):
            raise ValueError(f"File {file_key} is not a PDF document")

        try:
            download_url = await self.get_file_download_link(file_key)
            logger.info(f"Downloading PDF file from URL: {download_url}")

            async with httpx.AsyncClient(timeout=60.0) as client:
                response = await client.get(download_url)
                response.raise_for_status()

                pdf_reader = PyPDF2.PdfReader(io.BytesIO(response.content))

                # Извлекаем весь текст из PDF
                full_text = []
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    if page_text.strip():
                        full_text.append(page_text.strip())

                extracted_text = "\n".join(full_text)
                logger.info(
                    f"Extracted {len(extracted_text)} characters from PDF: {file_key} ({len(pdf_reader.pages)} pages)"
                )

                return extracted_text

        except Exception as e:
            logger.error(f"Error extracting text from PDF {file_key}: {e}")
            raise Exception(f"Не удалось извлечь текст из PDF файла: {str(e)}")

    @staticmethod
    def _normalize_prefix(prefix: str) -> str:
        normalized = prefix.lstrip("/")
        if normalized and not normalized.endswith("/"):
            normalized += "/"
        return normalized

    def _normalize_key(self, s3_key: str) -> str:
        normalized = s3_key.lstrip("/")
        if self.temp_prefix and not normalized.startswith(self.temp_prefix):
            normalized = f"{self.temp_prefix}{normalized}"
        return normalized
